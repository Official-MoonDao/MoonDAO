#!/usr/bin/env python3
"""Builds the 10-minute town hall deck for the MDP-249 Executive Branch final report.

Run:  python3 build/generate_deck.py
Output: dist/EB_Final_Report_Q2_Q3_2026_Townhall.pptx (convert to PDF with LibreOffice;
see README.md).

Every figure comes from docs/reports/eb-mdp-249 (report, charts and data/metrics.json),
so the slides and the written report stay in step.
"""
import os
import sys

from PIL import Image, ImageEnhance
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, '..'))
REPO = os.path.normpath(os.path.join(ROOT, '..', '..'))
RIO = os.path.join(REPO, 'presentations', 'rio-innovation-week-2026')
sys.path.insert(0, os.path.join(RIO, 'build'))

from deckutil import (  # noqa: E402
    SLIDE_W, SLIDE_H, MARGIN, CONTENT_W, FONT, FONT_HEAD,
    NAVY_DARK, NAVY, BLUE, ORANGE, RED, WHITE, BG_PANEL,
    LINE_GRAY, TEXT_DARK, TEXT_GRAY, TEXT_MUTE,
    set_bg, add_rect, add_line, add_text, add_rich, add_bullets,
    header, card, accent_rail, evenly_spaced,
)

CHARTS = os.path.join(REPO, 'docs', 'reports', 'eb-mdp-249', 'charts')
LOGO = os.path.join(RIO, 'assets', 'MoonDAO_icon.png')
CACHE = os.path.join(HERE, '_cache')
DIST = os.path.join(ROOT, 'dist')
OUT = os.path.join(DIST, 'EB_Final_Report_Q2_Q3_2026_Townhall.pptx')
os.makedirs(CACHE, exist_ok=True)
os.makedirs(DIST, exist_ok=True)

GREEN = RGBColor(0x2E, 0x8B, 0x57)
SOFT = RGBColor(0xC7, 0xCC, 0xE4)
AMBER = ORANGE
TOTAL = 13

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def chart(name):
    return os.path.join(CHARTS, name)


def new_slide(notes):
    s = prs.slides.add_slide(BLANK)
    set_bg(s, WHITE)
    s.notes_slide.notes_text_frame.text = notes
    return s


def footer(s, n):
    add_line(s, MARGIN, SLIDE_H - Inches(0.52), CONTENT_W, 0, color=LINE_GRAY, weight=Pt(0.75))
    add_text(s, MARGIN, SLIDE_H - Inches(0.42), Inches(9.5), Inches(0.28),
             "MoonDAO Town Hall  ·  Executive Branch Final Report, Q2 – Q3 2026 (MDP-249)  ·  "
             "Data as of 7 Oct 2026",
             size=Pt(9), color=TEXT_MUTE)
    add_text(s, SLIDE_W - MARGIN - Inches(1.3), SLIDE_H - Inches(0.42), Inches(1.3), Inches(0.28),
             f"{n:02d}  /  {TOTAL}", size=Pt(9), color=TEXT_MUTE, font=FONT_HEAD, align=PP_ALIGN.RIGHT)


def picture_fit(s, path, l, t, w, h, align='center'):
    """Place an image inside a box without distorting it."""
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    x = l + (w - pw) // 2 if align == 'center' else l
    y = t + (h - ph) // 2
    return s.shapes.add_picture(path, x, y, width=pw, height=ph)


def darkened(path, factor, ratio):
    out = os.path.join(CACHE, f"dark_{os.path.basename(path)}_{factor}_{ratio[0]:.2f}.png")
    if not os.path.exists(out):
        im = Image.open(path).convert('RGB')
        iw, ih = im.size
        tw = min(iw, int(ih * ratio[0] / ratio[1]))
        th = int(tw * ratio[1] / ratio[0])
        x, y = (iw - tw) // 2, (ih - th) // 2
        im = im.crop((x, y, x + tw, y + th))
        ImageEnhance.Brightness(im).enhance(factor).save(out)
    return out


def big_stat(s, l, t, w, h, number, label, sub=None, accent=NAVY):
    card(s, l, t, w, h)
    accent_rail(s, l, t, h, accent=accent)
    add_text(s, l + Inches(0.28), t + Inches(0.2), w - Inches(0.45), Inches(0.6), number,
             size=Pt(30), color=NAVY_DARK, bold=True, font=FONT_HEAD)
    add_text(s, l + Inches(0.28), t + Inches(0.9), w - Inches(0.45), Inches(0.35), label,
             size=Pt(13), color=TEXT_DARK, bold=True, font=FONT_HEAD)
    if sub:
        add_text(s, l + Inches(0.28), t + Inches(1.25), w - Inches(0.45), h - Inches(1.3), sub,
                 size=Pt(11.5), color=TEXT_GRAY)


def pill(s, l, t, w, text, color):
    add_rect(s, l, t, w, Inches(0.36), fill=color, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    add_text(s, l, t + Inches(0.06), w, Inches(0.3), text, size=Pt(11), color=WHITE,
             bold=True, font=FONT_HEAD, align=PP_ALIGN.CENTER)


# --------------------------------------------------------------------- 01 --
def slide_01():
    s = new_slide(
        "[0:20] Hi everyone. This is the Executive Branch's final report for MDP-249, the "
        "Q2 to Q3 2026 term, which ran from May through early September. I'm presenting on "
        "behalf of Pablo, Ryan and Miguel. Ten minutes: what we said we'd do, what happened, "
        "what it cost, and how we'd grade ourselves. All numbers are measured from on-chain "
        "data as of yesterday, October 7, and the full written report is posted with links "
        "to every transaction.")
    panel_x = Inches(8.4)
    panel_w = Emu(SLIDE_W - panel_x)
    bg = darkened(os.path.join(RIO, 'assets', 'moon-full.jpg'), 0.55,
                  (panel_w.inches, SLIDE_H.inches))
    s.shapes.add_picture(bg, panel_x, 0, width=panel_w, height=SLIDE_H)
    add_rect(s, panel_x, 0, Inches(0.055), SLIDE_H, fill=ORANGE)

    lx, lw = Inches(0.85), Inches(7.2)
    logo_w = Inches(2.6)
    s.shapes.add_picture(LOGO, lx, Inches(0.75), width=logo_w, height=int(logo_w * 345 / 1233))
    add_text(s, lx, Inches(2.0), lw, Inches(0.35), "MOONDAO TOWN HALL  ·  8 OCTOBER 2026",
             size=Pt(13), color=ORANGE, bold=True, font=FONT_HEAD)
    add_text(s, lx - Inches(0.03), Inches(2.4), lw, Inches(1.6),
             "Executive Branch\nFinal Report", size=Pt(46), color=NAVY_DARK, bold=True,
             font=FONT_HEAD, spacing=0.95)
    add_text(s, lx, Inches(4.05), lw, Inches(0.45), "Q2 2026 – Q3 2026  ·  MDP-249  ·  Project #131",
             size=Pt(20), color=NAVY, font=FONT_HEAD)
    add_line(s, lx, Inches(4.8), Inches(6.4), 0, color=LINE_GRAY, weight=Pt(0.75))
    team = [("Pablo Moncada-Larrotiz", "Executive Lead"),
            ("Ryan", "Communications, community & partnerships"),
            ("Miguel", "Frontend & product engineering")]
    for i, (name, role) in enumerate(team):
        y = Inches(5.05) + Inches(0.45) * i
        add_rich(s, lx, y, lw, Inches(0.4), [
            (name, Pt(15), NAVY_DARK, True, False, FONT_HEAD),
            (f"   {role}", Pt(13), TEXT_GRAY, False, False, FONT)])
    add_text(s, lx, Inches(6.75), lw, Inches(0.3),
             "Grades shown are provisional, for Executive Lead review",
             size=Pt(11), color=TEXT_MUTE, italic=True)


# --------------------------------------------------------------------- 02 --
def slide_02():
    s = new_slide(
        "[0:40] Quick reminder of the mandate. MDP-249 passed the Member House on May 26 with "
        "91% from 21 voters. Five months, $132,000 core budget plus a $24,000 bonus pool that "
        "only pays on verified milestones. Three objectives: first, put Frank in space or deploy "
        "DePrize with a lunar challenge. Second, grow the Space Acceleration Network and get "
        "Teams using jobs and the marketplace. Third, run the DAO's operations within budget, "
        "with a new project system, an audit, a for-profit plan and an election.")
    header(s, "What we committed to", "MDP-249: three objectives, five months, $132k", 2, LOGO)
    objs = [
        ("1", "Frank & DePrize",
         "Secure a seat for Frank, or deploy DePrize with the Frank White capital, and launch "
         "a competitive lunar simulation and prototyping initiative.", ORANGE),
        ("2", "Grow the network",
         "Citizens 199 → 300, Teams 20 → 30, 60% of Teams using jobs or the marketplace, "
         "discovery calls with 15+ Teams.", BLUE),
        ("3", "Run the DAO well",
         "Budget discipline, new project system, operational audit, for-profit plan, Realistic "
         "Goals, GDPR, an Executive Branch election.", NAVY),
    ]
    gap = Inches(0.35)
    cw = evenly_spaced(3, CONTENT_W, gap)
    top, ch = Inches(1.65), Inches(3.3)
    for i, (num, title, body, accent) in enumerate(objs):
        x = MARGIN + (cw + gap) * i
        card(s, x, top, cw, ch)
        add_rect(s, x, top, cw, Inches(0.08), fill=accent)
        add_text(s, x + Inches(0.35), top + Inches(0.35), Inches(1), Inches(0.7), num,
                 size=Pt(40), color=accent, bold=True, font=FONT_HEAD)
        add_text(s, x + Inches(0.35), top + Inches(1.15), cw - Inches(0.7), Inches(0.45), title,
                 size=Pt(19), color=NAVY_DARK, bold=True, font=FONT_HEAD)
        add_text(s, x + Inches(0.35), top + Inches(1.7), cw - Inches(0.7), Inches(1.5), body,
                 size=Pt(13.5), color=TEXT_GRAY, spacing=1.1)
    facts = [("Passed 26 May", "91% · 21 voters"), ("Funded", "1 May – 3 Sep 2026"),
             ("Core budget", "$132,000"), ("Bonus pool", "$24,000 at risk")]
    fw = evenly_spaced(4, CONTENT_W, gap)
    for i, (k, v) in enumerate(facts):
        x = MARGIN + (fw + gap) * i
        card(s, x, Inches(5.3), fw, Inches(1.1), fill=BG_PANEL, shadow=False)
        add_text(s, x + Inches(0.28), Inches(5.45), fw, Inches(0.3), k.upper(),
                 size=Pt(10.5), color=ORANGE, bold=True, font=FONT_HEAD)
        add_text(s, x + Inches(0.28), Inches(5.78), fw - Inches(0.4), Inches(0.5), v,
                 size=Pt(16), color=NAVY_DARK, bold=True, font=FONT_HEAD)
    footer(s, 2)


# --------------------------------------------------------------------- 03 --
def slide_03():
    s = new_slide(
        "[0:50] The headline numbers. DePrize is live on Arbitrum with five competitions. "
        "Citizens grew 33%, from 199 to 265 minted, 238 currently active, the fastest growth "
        "we've had, but short of 300. Teams went from 20 to 25. Spending landed within 0.6% of "
        "the budget. And the three of us merged 359 pull requests, roughly double the previous "
        "pace, largely thanks to AI-assisted development.")
    header(s, "At a glance", "The term in six numbers", 3, LOGO)
    stats = [
        ("Live", "DePrize on Arbitrum", "Mainnet 25 Aug · 5 competitions on mainnet, 8 on Sepolia", ORANGE),
        ("+33%", "Citizens", "199 → 265 minted · 238 active (target 300)", BLUE),
        ("20 → 25", "Teams", "Five new Teams, all active (target 30)", BLUE),
        ("44%", "Teams using services", "11 of 25 posted a job or listing (target 60%)", BLUE),
        ("$132.7k", "Spent vs $132k budget", "Within 0.6% · payroll via LlamaPay streams", NAVY),
        ("359", "Merged pull requests", "Three-person team · 1 May – 7 Oct", NAVY),
    ]
    gap = Inches(0.32)
    cw = evenly_spaced(3, CONTENT_W, gap)
    ch = Inches(2.3)
    for i, (n, label, sub, accent) in enumerate(stats):
        r, c = divmod(i, 3)
        big_stat(s, MARGIN + (cw + gap) * c, Inches(1.65) + (ch + gap) * r, cw, ch, n, label, sub, accent)
    footer(s, 3)


# --------------------------------------------------------------------- 04 --
def slide_04():
    s = new_slide(
        "[1:00] Objective 1. Phase 1 of the Overview Flight raised capital, but you voted to "
        "keep fundraising rather than convert it into a prize pool, so we built DePrize as "
        "general MoonDAO infrastructure instead. It went live on Arbitrum on August 25 with a "
        "demonstration prize, The Moon Is A Harsh Mistress. A security review found a bug in "
        "the market contract, which we fixed and redeployed before any race went live. Then "
        "on September 26 we registered four lunar capability races: Touchdown, Night Shift, "
        "First Tracks and Water Ice. We ran the full lifecycle on testnet, including declaring "
        "a winner and paying out. Honest caveats: this landed about eight weeks after the "
        "month-2 target, and so far the only bets are our own test bets, so no revenue yet.")
    header(s, "Objective 1  ·  DePrize", "DePrize is live on Arbitrum", 4, LOGO)
    events = [
        ("25 Aug", "Mainnet launch", "DePrize #1: The Moon Is A Harsh Mistress (demo prize)"),
        ("11 Sep", "Security fix", "Bug H-01 found in review, fixed and redeployed before any race"),
        ("18 Sep", "Full lifecycle", "Sepolia: bet, declare winner, redeem, end to end"),
        ("26 Sep", "Four lunar races", "Registered on Arbitrum, each with its own market"),
    ]
    ty = Inches(2.05)
    add_line(s, MARGIN + Inches(0.2), ty, CONTENT_W - Inches(0.4), 0, color=NAVY, weight=Pt(2.5))
    ew = CONTENT_W / 4
    for i, (date, title, body) in enumerate(events):
        x = MARGIN + ew * i
        d = Inches(0.26)
        add_rect(s, x + Inches(0.2), ty - d / 2, d, d, fill=ORANGE, shape_type=MSO_SHAPE.OVAL)
        add_text(s, x + Inches(0.2), Inches(1.5), ew, Inches(0.3), date.upper(),
                 size=Pt(12), color=ORANGE, bold=True, font=FONT_HEAD)
        add_text(s, x + Inches(0.2), ty + Inches(0.3), ew - Inches(0.4), Inches(0.35), title,
                 size=Pt(15), color=NAVY_DARK, bold=True, font=FONT_HEAD)
        add_text(s, x + Inches(0.2), ty + Inches(0.68), ew - Inches(0.45), Inches(0.8), body,
                 size=Pt(12), color=TEXT_GRAY, spacing=1.05)
    races = [("Touchdown", "Next upright working lunar landing", "6 competitors"),
             ("Night Shift", "First machine to work through a lunar night", "8 competitors"),
             ("First Tracks", "Commercial rover egress and drive", "6 competitors"),
             ("Water Ice", "First in-situ surface water ice", "4 competitors")]
    gap = Inches(0.25)
    rw = evenly_spaced(4, CONTENT_W, gap)
    ry, rh = Inches(3.85), Inches(1.5)
    for i, (name, what, n) in enumerate(races):
        x = MARGIN + (rw + gap) * i
        card(s, x, ry, rw, rh)
        accent_rail(s, x, ry, rh, accent=BLUE)
        add_text(s, x + Inches(0.28), ry + Inches(0.18), rw - Inches(0.4), Inches(0.35), name,
                 size=Pt(16), color=NAVY_DARK, bold=True, font=FONT_HEAD)
        add_text(s, x + Inches(0.28), ry + Inches(0.58), rw - Inches(0.45), Inches(0.6), what,
                 size=Pt(12), color=TEXT_GRAY, spacing=1.05)
        add_text(s, x + Inches(0.28), ry + rh - Inches(0.38), rw - Inches(0.4), Inches(0.3), n,
                 size=Pt(11), color=BLUE, bold=True, font=FONT_HEAD)
    card(s, MARGIN, Inches(5.55), CONTENT_W, Inches(0.85), fill=BG_PANEL, shadow=False)
    add_rich(s, MARGIN + Inches(0.3), Inches(5.78), CONTENT_W - Inches(0.6), Inches(0.45), [
        ("Not yet:  ", Pt(14), RED, True, False, FONT_HEAD),
        ("public betting on the races, an analog prototype test, and revenue. Landed ~8 weeks "
         "after the month-2 target; mainnet bets so far are internal tests (0.0127 ETH).",
         Pt(13.5), TEXT_DARK, False, False, FONT)])
    footer(s, 4)


# --------------------------------------------------------------------- 05 --
def slide_05():
    s = new_slide(
        "[0:50] Instead of running the lunar simulation as an outside challenge, Miguel built "
        "it in-house. This is Moon Base Zero: a true-to-scale base on the Shackleton connecting "
        "ridge, with real sun position and terrain. It models 50 projects from 33 organizations "
        "across 12 shared goals, and every competitor in every DePrize race has its own lot. On "
        "the right is the Touchdown race: the six landers, the rules, and live odds read from "
        "the market. It's behind an access gate for now. [Optional: switch to the live demo.]")
    header(s, "Objective 1  ·  Lunar simulation", "Moon Base Zero: every race, on the Moon", 5, LOGO)
    gap = Inches(0.3)
    w = evenly_spaced(2, CONTENT_W, gap)
    h = int(w * 9 / 16)
    top = Inches(1.6)
    for i, (img, cap) in enumerate([
            ('14-moon-base-zero.png', "The base on the Shackleton connecting ridge, with the four live races"),
            ('15-moon-base-zero-touchdown.png', "Touchdown race: six landers, rules and live DePrize odds")]):
        x = MARGIN + (w + gap) * i
        add_rect(s, x - Inches(0.02), top - Inches(0.02), w + Inches(0.04), h + Inches(0.04), fill=LINE_GRAY)
        s.shapes.add_picture(chart(img), x, top, width=w, height=h)
        add_text(s, x, top + h + Inches(0.12), w, Inches(0.3), cap, size=Pt(11.5), color=TEXT_GRAY)
    stats = ["50 projects", "33 organizations", "12 capability goals", "336 unit tests", "First shipped 23 Jul"]
    sw = evenly_spaced(5, CONTENT_W, Inches(0.2))
    for i, t in enumerate(stats):
        x = MARGIN + (sw + Inches(0.2)) * i
        card(s, x, Inches(5.75), sw, Inches(0.6), fill=BG_PANEL, shadow=False)
        add_text(s, x, Inches(5.9), sw, Inches(0.35), t, size=Pt(14), color=NAVY_DARK, bold=True,
                 font=FONT_HEAD, align=PP_ALIGN.CENTER)
    footer(s, 5)


# --------------------------------------------------------------------- 06 --
def slide_06():
    s = new_slide(
        "[0:50] On Frank's flight: after Phase 1 closed in April we approached 12 carriers and "
        "brought the findings to you. In the June path vote, 92.9% of OVERVIEW voting weight "
        "chose to keep options open and keep fundraising. The campaign reopened on July 9 "
        "with a single $250k goal. Lifetime it's raised 27.4 ETH from 180 contributions, but "
        "re-engaging after a pause is hard: only 0.64 ETH came in after the reopen. "
        "[Only if ready to share: agreements are in place for Frank's flight; an announcement "
        "is coming.]")
    header(s, "Objective 1  ·  Overview Flight", "You chose to keep fundraising for Frank", 6, LOGO)
    lw = Inches(5.6)
    add_text(s, MARGIN, Inches(1.55), lw, Inches(0.3), "PATH VOTE  ·  15–22 JUNE  ·  31 VOTERS",
             size=Pt(11), color=ORANGE, bold=True, font=FONT_HEAD)
    options = [("B · Keep options open, keep fundraising", 92.9, 22, BLUE),
               ("A · Commit now to a stratospheric balloon", 0.2, 3, TEXT_MUTE),
               ("C · Refund contributors", 0.7, 2, TEXT_MUTE),
               ("Abstain", 6.2, 4, TEXT_MUTE)]
    bar_w = lw - Inches(1.1)
    for i, (label, pct, voters, color) in enumerate(options):
        y = Inches(2.0) + Inches(0.78) * i
        add_text(s, MARGIN, y, lw, Inches(0.3), f"{label}  ({voters} voters)",
                 size=Pt(12.5), color=TEXT_DARK, bold=(i == 0), font=FONT_HEAD if i == 0 else FONT)
        add_rect(s, MARGIN, y + Inches(0.32), bar_w, Inches(0.24), fill=BG_PANEL)
        add_rect(s, MARGIN, y + Inches(0.32), max(int(bar_w * pct / 100), Inches(0.04)), Inches(0.24), fill=color)
        add_text(s, MARGIN + bar_w + Inches(0.12), y + Inches(0.28), Inches(1.0), Inches(0.3), f"{pct}%",
                 size=Pt(14), color=NAVY_DARK, bold=True, font=FONT_HEAD)
    add_text(s, MARGIN, Inches(5.2), lw, Inches(0.3), "Share of 10,444 $OVERVIEW voting weight",
             size=Pt(10.5), color=TEXT_MUTE)
    rx = MARGIN + lw + Inches(0.4)
    rw = CONTENT_W - lw - Inches(0.4)
    picture_fit(s, chart('10-overview-flight-contributions.png'), rx, Inches(1.5), rw, Inches(3.4))
    sw = evenly_spaced(3, CONTENT_W, Inches(0.3))
    for i, (n, label, sub) in enumerate([
            ("27.39 ETH", "Raised in total", "180 contributions"),
            ("0.64 ETH", "Since reopening", "9 Jul onward · 21 contributions · 16 wallets"),
            ("12", "Carriers approached", "Balloon, suborbital, orbital")]):
        x = MARGIN + (sw + Inches(0.3)) * i
        card(s, x, Inches(5.6), sw, Inches(0.85), fill=BG_PANEL, shadow=False)
        add_rich(s, x + Inches(0.3), Inches(5.72), sw - Inches(0.5), Inches(0.35), [
            (n, Pt(17), NAVY_DARK, True, False, FONT_HEAD),
            (f"   {label}", Pt(12), TEXT_DARK, True, False, FONT_HEAD)])
        add_text(s, x + Inches(0.3), Inches(6.08), sw - Inches(0.5), Inches(0.3), sub,
                 size=Pt(11), color=TEXT_GRAY)
    footer(s, 6)


# --------------------------------------------------------------------- 07 --
def slide_07():
    s = new_slide(
        "[0:50] Objective 2, the network. Citizens grew from 199 to 265 minted, 66 new Citizens "
        "in five months versus 35 in the five before. 238 are currently active, because 27 have "
        "lapsed, which is why we shipped renewal prompts. Teams grew from 20 to 25: Habitat "
        "Marte, A Heart for Space, the U.S. Space and Rocket Center Education Foundation, "
        "Geração de Marte and Zephalto. Both missed their targets of 300 and 30. Those targets "
        "assumed a live mid-cycle raise, which was our biggest acquisition engine last time.")
    header(s, "Objective 2  ·  Network growth", "Fastest growth on record, short of target", 7, LOGO)
    picture_fit(s, chart('02-citizen-growth.png'), MARGIN, Inches(1.45), Inches(7.9), Inches(5.0))
    rx = MARGIN + Inches(8.15)
    rw = CONTENT_W - Inches(8.15)
    big_stat(s, rx, Inches(1.6), rw, Inches(1.6), "199 → 265", "Citizens minted",
             "238 active · target 300", BLUE)
    big_stat(s, rx, Inches(3.4), rw, Inches(1.6), "20 → 25", "Teams", "Target 30", BLUE)
    add_text(s, rx, Inches(5.2), rw, Inches(1.3),
             "New Teams: Habitat Marte, A Heart for Space, USSRC Education Foundation, "
             "Geração de Marte, Zephalto",
             size=Pt(12), color=TEXT_GRAY, spacing=1.1)
    footer(s, 7)


# --------------------------------------------------------------------- 08 --
def slide_08():
    s = new_slide(
        "[0:40] Utilization. 11 of our 25 Teams, 44%, used at least one service this term "
        "against a 60% target. The marketplace beat its goal with 17 new listings, mostly real "
        "space experiences: analog missions, zero-g flights, training, MDRS seats. Jobs fell "
        "short at 7 new listings. Job posting is concentrated in a few Teams, so the board "
        "needs seeding, not just availability.")
    header(s, "Objective 2  ·  Utilization", "Marketplace beat its target; jobs did not", 8, LOGO)
    picture_fit(s, chart('05-team-utilization.png'), MARGIN, Inches(1.45), Inches(7.9), Inches(5.0))
    rx = MARGIN + Inches(8.15)
    rw = CONTENT_W - Inches(8.15)
    big_stat(s, rx, Inches(1.6), rw, Inches(1.5), "17 / 15", "Marketplace listings", "Target met", GREEN)
    big_stat(s, rx, Inches(3.3), rw, Inches(1.5), "7 / 15", "Job listings", "Target missed", RED)
    big_stat(s, rx, Inches(5.0), rw, Inches(1.5), "44%", "Teams using a service", "6 jobs · 8 marketplace · 3 both", AMBER)
    footer(s, 8)


# --------------------------------------------------------------------- 09 --
def slide_09():
    s = new_slide(
        "[0:50] Ryan led about 30 discovery calls with 21 Teams between late May and "
        "September, a little over two a week: 11 existing Teams and 10 prospective ones. Two "
        "themes came up again and again, and we shipped both. First, Teams wanted to post "
        "detailed roles and share them, so listings now show full descriptions and have their "
        "own shareable links, and Team pages are visible to everyone. Second, buyers struggled "
        "with crypto and international payments, so the marketplace now has a card and bank "
        "onramp. On the prospective side, Lonestar has tentatively agreed and Stardust is "
        "whitelisted.")
    header(s, "Objective 2  ·  Discovery", "30 calls with 21 Teams, two fixes shipped", 9, LOGO)
    lw = Inches(4.3)
    big_stat(s, MARGIN, Inches(1.6), lw, Inches(1.55), "~30 calls", "with 21 Teams",
             "Late May – September · ~2 a week", BLUE)
    card(s, MARGIN, Inches(3.35), lw, Inches(3.05), fill=BG_PANEL, shadow=False)
    add_text(s, MARGIN + Inches(0.3), Inches(3.5), lw - Inches(0.5), Inches(0.3), "11 IN THE NETWORK",
             size=Pt(10.5), color=ORANGE, bold=True, font=FONT_HEAD)
    add_text(s, MARGIN + Inches(0.3), Inches(3.8), lw - Inches(0.5), Inches(1.1),
             "LifeShip, LunARC, The Inspired 24, Mars Society, The Human Space Program, Celestial "
             "Commons, SpaceKind, Habitat Marte, Space Camp USA, Generation Mars Brazil, Zephalto",
             size=Pt(11), color=TEXT_DARK, spacing=1.05)
    add_text(s, MARGIN + Inches(0.3), Inches(5.0), lw - Inches(0.5), Inches(0.3), "10 PROSPECTIVE",
             size=Pt(10.5), color=ORANGE, bold=True, font=FONT_HEAD)
    add_text(s, MARGIN + Inches(0.3), Inches(5.3), lw - Inches(0.5), Inches(1.0),
             "Including Lonestar (tentatively agreed) and Stardust (whitelisted)",
             size=Pt(11), color=TEXT_DARK, spacing=1.05)

    rx = MARGIN + lw + Inches(0.35)
    rw = CONTENT_W - lw - Inches(0.35)
    items = [
        ("1", "Richer, shareable listings",
         "Full descriptions instead of truncated text, a deep link for every listing, and Team "
         "pages visible to non-Citizens so links work for anyone.",
         "Asked for by The Inspired 24, Habitat Marte, LunCO"),
        ("2", "Card & bank onramp for purchases",
         "Buyers fund a wallet and buy a listing in one flow without holding crypto. Onboarding "
         "no longer steers people to a single provider.",
         "Asked for by Habitat Marte (international payments), LunARC (crypto hurdles)"),
    ]
    ih = Inches(2.3)
    for i, (n, title, body, who) in enumerate(items):
        y = Inches(1.6) + (ih + Inches(0.2)) * i
        card(s, rx, y, rw, ih)
        accent_rail(s, rx, y, ih, accent=ORANGE)
        add_text(s, rx + Inches(0.35), y + Inches(0.22), rw - Inches(0.6), Inches(0.4),
                 f"Shipped  {n}   ·   {title}", size=Pt(16), color=NAVY_DARK, bold=True, font=FONT_HEAD)
        add_text(s, rx + Inches(0.35), y + Inches(0.75), rw - Inches(0.6), Inches(0.95), body,
                 size=Pt(13), color=TEXT_GRAY, spacing=1.1)
        add_text(s, rx + Inches(0.35), y + ih - Inches(0.5), rw - Inches(0.6), Inches(0.3), who,
                 size=Pt(11), color=BLUE, italic=True)
    footer(s, 9)


# --------------------------------------------------------------------- 10 --
def slide_10():
    s = new_slide(
        "[0:50] Objective 3, operations and governance. Done: the new project system with 100% "
        "of initiatives migrated, later refined into Project System v9; the operational audit, "
        "published as the August financial disclosure plus a live internal dashboard; GDPR, "
        "which opened Europe to browsing while only Citizen creation is restricted; and the "
        "Realistic Goals sequencing, which became the objectives in the next EB proposal, "
        "MDP-271. Partly done: the for-profit plan is drafted but not yet presented. Not done: "
        "no cost reduction, because AI tooling costs rose, and the Executive Lead election vote "
        "was postponed; we'll announce the new date.")
    header(s, "Objective 3  ·  Operations & governance", "Most of the operating agenda delivered", 10, LOGO)
    rows = [
        ("Done", GREEN, [
            ("Project system: ", "100% migrated, then Project System v9 (MDP-267)"),
            ("Operational audit: ", "14 Aug financial disclosure and burn report, live dashboard"),
            ("GDPR: ", "Europe opened to browsing; only Citizen creation restricted"),
            ("Realistic Goals: ", "delivered as the OKRs of the next EB proposal, MDP-271"),
        ]),
        ("Partly", AMBER, [
            ("For-profit arm: ", "plan and business case drafted, not yet presented"),
            ("Burn reporting: ", "continuous internal dashboard instead of monthly reports"),
        ]),
        ("Not yet", RED, [
            ("Cost reduction: ", "none; AI-tooling costs rose"),
            ("Election: ", "nominations and town hall held; Member House vote postponed"),
        ]),
    ]
    y = Inches(1.6)
    for label, color, items in rows:
        h = Inches(0.42) * len(items) + Inches(0.35)
        card(s, MARGIN, y, CONTENT_W, h)
        accent_rail(s, MARGIN, y, h, accent=color)
        pill(s, MARGIN + Inches(0.3), y + (h - Inches(0.36)) / 2, Inches(1.2), label, color)
        add_bullets(s, MARGIN + Inches(1.8), y + Inches(0.2), CONTENT_W - Inches(2.1), h - Inches(0.3),
                    items, size=Pt(14), gap=Pt(4), marker='·', bullet_color=color)
        y += h + Inches(0.18)
    footer(s, 10)


# --------------------------------------------------------------------- 11 --
def slide_11():
    s = new_slide(
        "[0:50] The money. We were funded once in USDC for May and once in ETH for the "
        "remaining four months. Payroll ran through LlamaPay vesting streams, which cut payroll "
        "transactions from about 12 to 2, though it did expose pay to ETH price moves. Total "
        "paid: $132,727 against $132,000, within 0.6%. No bonus has been paid. On revenue: the "
        "ETH that reaches the treasury on-chain from Citizen and Team subscriptions and "
        "Launchpad fees is about 1.67 ETH, roughly $4,300, over the last twelve months. Card "
        "and fiat checkouts aren't in that number. Getting to cash-flow sustainability is the "
        "core of the next proposal.")
    header(s, "Budget & treasury", "On budget: $132.7k spent of $132k", 11, LOGO)
    picture_fit(s, chart('07-budget-vs-actual.png'), MARGIN, Inches(1.45), Inches(7.9), Inches(4.9))
    rx = MARGIN + Inches(8.15)
    rw = CONTENT_W - Inches(8.15)
    big_stat(s, rx, Inches(1.6), rw, Inches(1.5), "+0.6%", "vs core budget",
             "$120.5k payroll + $12.2k ops & flex", NAVY)
    big_stat(s, rx, Inches(3.3), rw, Inches(1.5), "$0", "Bonus paid", "Of the $24k at-risk pool", NAVY)
    big_stat(s, rx, Inches(5.0), rw, Inches(1.5), "≈ $4.3k", "On-chain revenue, 12 months",
             "1.67 ETH to the treasury; excludes card/fiat", AMBER)
    footer(s, 11)


# --------------------------------------------------------------------- 12 --
def slide_12():
    s = new_slide(
        "[0:40] How we'd grade ourselves, provisionally. Objective 1, Meets: the infrastructure "
        "went beyond scope, but it was late and there's no revenue yet. Objective 2, Does Not "
        "Meet: real growth and the discovery calls done, but all three network targets missed. "
        "Objective 3, Meets, borderline: operations, budget, the project system and GDPR "
        "delivered; the election and cost reduction weren't. On bonuses, nothing has been paid; "
        "only the project-system milestone is a clean claim, and anything beyond that is for "
        "the Executive Leads and Senate to decide.")
    header(s, "Self-assessment", "Provisional grades", 12, LOGO)
    grades = [
        ("1", "Frank & DePrize", "Meets", GREEN,
         "Beyond scope on infrastructure; late; no revenue or prototype test yet"),
        ("2", "Network growth", "Does not meet", RED,
         "Strong growth and discovery done; Citizen, Team and utilization targets missed"),
        ("3", "Operations", "Meets (borderline)", AMBER,
         "Budget, project system, audit, GDPR done; election and cost cuts not"),
    ]
    gap = Inches(0.35)
    cw = evenly_spaced(3, CONTENT_W, gap)
    top, ch = Inches(1.6), Inches(3.0)
    for i, (n, title, grade, color, why) in enumerate(grades):
        x = MARGIN + (cw + gap) * i
        card(s, x, top, cw, ch)
        add_rect(s, x, top, cw, Inches(0.08), fill=color)
        add_text(s, x + Inches(0.35), top + Inches(0.3), cw - Inches(0.7), Inches(0.35),
                 f"OBJECTIVE {n}", size=Pt(11), color=ORANGE, bold=True, font=FONT_HEAD)
        add_text(s, x + Inches(0.35), top + Inches(0.62), cw - Inches(0.7), Inches(0.4), title,
                 size=Pt(17), color=NAVY_DARK, bold=True, font=FONT_HEAD)
        pill(s, x + Inches(0.35), top + Inches(1.2), cw - Inches(0.7), grade, color)
        add_text(s, x + Inches(0.35), top + Inches(1.8), cw - Inches(0.7), Inches(1.1), why,
                 size=Pt(13), color=TEXT_GRAY, spacing=1.1)
    card(s, MARGIN, Inches(4.85), CONTENT_W, Inches(1.5), fill=BG_PANEL, shadow=False)
    add_text(s, MARGIN + Inches(0.35), Inches(5.0), Inches(3), Inches(0.3), "PERFORMANCE BONUS",
             size=Pt(11), color=ORANGE, bold=True, font=FONT_HEAD)
    add_text(s, MARGIN + Inches(0.35), Inches(5.35), CONTENT_W - Inches(0.7), Inches(0.9),
             "$0 of the $24,000 pool paid. Of the 10 claimable items, only \"project system 100% "
             "live\" ($1,000) is clearly met. DePrize launch timing and the lunar challenge are "
             "judgment calls for the Executive Leads and the Senate.",
             size=Pt(14), color=TEXT_DARK, spacing=1.1)
    footer(s, 12)


# --------------------------------------------------------------------- 13 --
def slide_13():
    s = new_slide(
        "[0:40] What we learned and where it goes next. Letting OVERVIEW holders decide "
        "Frank's path was right, but it meant DePrize launched without a seed pool. Growth "
        "follows activity, so the next term needs a live raise. Teams told us payments were the "
        "barrier, and we fixed it. And AI moved cost from salaries to tooling, so we should "
        "measure cost per outcome. All of this feeds MDP-271, the next EB proposal. The full "
        "report with every transaction link is on the project page. Happy to take questions.")
    set_bg(s, NAVY_DARK)
    add_rect(s, 0, 0, SLIDE_W, Inches(0.08), fill=ORANGE)
    add_text(s, MARGIN, Inches(0.55), Inches(9), Inches(0.3), "LEARNINGS  →  NEXT TERM",
             size=Pt(12), color=ORANGE, bold=True, font=FONT_HEAD)
    add_text(s, MARGIN, Inches(0.9), Inches(11), Inches(0.7), "What we're carrying into MDP-271",
             size=Pt(32), color=WHITE, bold=True, font=FONT_HEAD)
    lessons = [
        ("Let holders decide. ", "The path vote gave a clear mandate, but DePrize launched without a seed pool."),
        ("Growth follows activity. ", "Our best months had a live campaign; the next term needs a raise."),
        ("Remove payment friction. ", "Teams named payments as the blocker; the onramp is now live."),
        ("Measure cost per outcome. ", "AI doubled throughput but moved spend from salaries to tooling."),
    ]
    add_bullets(s, MARGIN, Inches(2.0), Inches(8.2), Inches(4.0), lessons, size=Pt(17),
                color=WHITE, gap=Pt(16), bullet_color=ORANGE)
    for p in s.shapes[-1].text_frame.paragraphs:
        for r in p.runs[2:]:
            r.font.color.rgb = SOFT
    px = MARGIN + Inches(8.8)
    pw = CONTENT_W - Inches(8.8)
    add_rect(s, px, Inches(2.0), pw, Inches(3.6), fill=NAVY, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    add_text(s, px + Inches(0.35), Inches(2.3), pw - Inches(0.7), Inches(0.5), "Questions?",
             size=Pt(28), color=WHITE, bold=True, font=FONT_HEAD)
    add_text(s, px + Inches(0.35), Inches(3.1), pw - Inches(0.7), Inches(2.3),
             "Full report, data and every transaction link:\nmoondao.com/project/131\n\n"
             "Next EB proposal:\nmoondao.com/project/153",
             size=Pt(14), color=SOFT, spacing=1.15)
    add_text(s, MARGIN, SLIDE_H - Inches(0.6), CONTENT_W, Inches(0.3),
             "Thank you to every Citizen, Team and contributor who built with us this term.",
             size=Pt(12), color=TEXT_MUTE, italic=True)


for fn in [slide_01, slide_02, slide_03, slide_04, slide_05, slide_06, slide_07,
           slide_08, slide_09, slide_10, slide_11, slide_12, slide_13]:
    fn()

prs.save(OUT)
print(f"wrote {OUT}")
